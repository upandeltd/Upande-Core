# Copyright (c) 2026, ghost-mann and contributors
# For license information, please see license.txt

import os

import frappe
import openpyxl
from frappe.tests import IntegrationTestCase

from role_advisor import capability, module_workbook


class TestModuleWorkbook(IntegrationTestCase):
	@classmethod
	def setUpClass(cls):
		super().setUpClass()
		capability.clear_capability_index()
		cls.data = module_workbook.gather()
		cls.path = module_workbook.build(
			path=os.path.join(
				frappe.utils.get_site_path("private", "files"), "_ra_test_modwb.xlsx"
			)
		)
		cls.book = openpyxl.load_workbook(cls.path, read_only=True)

	@classmethod
	def tearDownClass(cls):
		cls.book.close()
		if os.path.exists(cls.path):
			os.remove(cls.path)
		super().tearDownClass()

	def test_there_is_a_sheet_for_every_module_with_grants(self):
		expected = len(self.data["by_module"])

		# Overview and Profile Index sit in front of the module tabs.
		self.assertEqual(len(self.book.sheetnames), expected + 2)
		self.assertEqual(self.book.sheetnames[0], "Overview")
		self.assertEqual(self.book.sheetnames[1], "Profile Index")

	def test_tab_names_are_unique_and_within_the_excel_limit(self):
		for name in self.book.sheetnames:
			self.assertLessEqual(len(name), 31, name)
			for char in "[]:*?/\\":
				self.assertNotIn(char, name)
		self.assertEqual(len(self.book.sheetnames), len(set(self.book.sheetnames)))

	def test_every_module_sheet_groups_roles_and_permissions_under_a_profile(self):
		module = max(self.data["by_module"], key=lambda m: len(self.data["by_module"][m]))
		rows = module_workbook.module_rows(module, self.data)
		sections = [row["section"] for row in rows]

		self.assertIn(module_workbook.SECTION_PROFILE, sections)
		self.assertIn(module_workbook.SECTION_PERMISSION, sections)
		# A PROFILE header must come before the rows it groups.
		self.assertEqual(sections[0], module_workbook.SECTION_PROFILE)

	def test_permission_rows_match_the_capability_index(self):
		"""A sheet that under-reports grants would understate real access."""
		for module, profiles in list(self.data["by_module"].items())[:6]:
			rows = module_workbook.module_rows(module, self.data)
			counted = sum(
				1 for r in rows if r["section"] == module_workbook.SECTION_PERMISSION
			)
			expected = sum(len(doctypes) for doctypes in profiles.values())
			self.assertEqual(counted, expected, module)

	def test_a_cross_module_profile_names_its_other_sheets(self):
		"""The reconciliation flag: owners must see who else edits this profile."""
		spread = max(self.data["profile_modules"], key=lambda p: len(self.data["profile_modules"][p]))
		module = sorted(self.data["profile_modules"][spread])[0]

		rows = module_workbook.module_rows(module, self.data)
		header = next(
			r for r in rows
			if r["section"] == module_workbook.SECTION_PROFILE and r["role_profile"] == spread
		)

		self.assertTrue(header["also_on_sheets"], f"{spread} spans many modules")
		self.assertNotIn(module, header["also_on_sheets"].split(", "))

	def test_add_rows_are_present_for_owner_input(self):
		module = sorted(self.data["by_module"])[0]
		rows = module_workbook.module_rows(module, self.data)

		self.assertIn(module_workbook.SECTION_ADD, [r["section"] for r in rows])

	def test_the_export_writes_nothing_to_the_site(self):
		before = (
			frappe.db.count("Role Profile"),
			frappe.db.count("Custom DocPerm"),
			frappe.db.count("Has Role"),
		)

		second = os.path.join(
			frappe.utils.get_site_path("private", "files"), "_ra_test_modwb2.xlsx"
		)
		module_workbook.build(path=second)
		os.remove(second)

		self.assertEqual(
			before,
			(
				frappe.db.count("Role Profile"),
				frappe.db.count("Custom DocPerm"),
				frappe.db.count("Has Role"),
			),
		)
