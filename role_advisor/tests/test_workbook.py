# Copyright (c) 2026, ghost-mann and contributors
# For license information, please see license.txt

import os

import frappe
import openpyxl
from frappe.tests import IntegrationTestCase

from role_advisor import capability, workbook


class TestWorkbook(IntegrationTestCase):
	@classmethod
	def setUpClass(cls):
		super().setUpClass()
		capability.clear_capability_index()
		cls.path = workbook.build(
			path=os.path.join(frappe.utils.get_site_path("private", "files"), "_ra_test_workbook.xlsx")
		)
		cls.book = openpyxl.load_workbook(cls.path, read_only=True)

	@classmethod
	def tearDownClass(cls):
		cls.book.close()
		if os.path.exists(cls.path):
			os.remove(cls.path)
		super().tearDownClass()

	def test_all_four_sheets_are_present_and_ordered(self):
		self.assertEqual(
			self.book.sheetnames, ["Users", "Role Profiles", "Roles", "Permissions"]
		)

	def test_row_counts_match_the_database(self):
		"""A sheet shorter than the table is a silently truncated export."""
		expected = {
			# Every user, enabled or not: a disabled account still holds roles
			# and a profile, and re-enabling it restores them silently.
			"Users": frappe.db.count("User"),
			"Role Profiles": frappe.db.count("Role Profile"),
			"Roles": frappe.db.count("Role"),
		}
		for title, count in expected.items():
			# +1 for the header row.
			self.assertEqual(self.book[title].max_row, count + 1, title)

	def test_permissions_sheet_has_one_row_per_profile_doctype_pair(self):
		index = capability.build_capability_index()
		pairs = sum(len(caps) for caps in index.values())

		self.assertEqual(self.book["Permissions"].max_row, pairs + 1)

	def test_editable_columns_are_labelled(self):
		"""The label is the only thing stopping a reviewer editing a fact column."""
		headers = [cell.value for cell in next(self.book["Role Profiles"].iter_rows(max_row=1))]

		self.assertIn("proposed_name (EDITABLE)", headers)
		self.assertIn("action (EDITABLE)", headers)
		self.assertIn("role_profile", headers)
		self.assertNotIn("role_profile (EDITABLE)", headers)

	def test_the_export_writes_nothing_to_the_site(self):
		before = (
			frappe.db.count("Role Profile"),
			frappe.db.count("Role"),
			frappe.db.count("User Role Profile"),
		)

		workbook.build(path=os.path.join(frappe.utils.get_site_path("private", "files"), "_ra_test_workbook2.xlsx"))

		self.assertEqual(
			before,
			(
				frappe.db.count("Role Profile"),
				frappe.db.count("Role"),
				frappe.db.count("User Role Profile"),
			),
		)
		os.remove(os.path.join(frappe.utils.get_site_path("private", "files"), "_ra_test_workbook2.xlsx"))

	def test_disabled_users_are_included(self):
		"""Filtering to enabled=1 is what under-reported the user list."""
		sheet = self.book["Users"]
		headers = [cell.value for cell in next(sheet.iter_rows(max_row=1))]
		column = headers.index("enabled")

		values = [row[column] for row in sheet.iter_rows(min_row=2, values_only=True)]

		self.assertIn(0, values, "expected at least one disabled user")
		self.assertEqual(len(values), frappe.db.count("User"))

	def test_duplicate_profiles_are_named_in_the_duplicate_column(self):
		"""Upande Team's three names grant identically; the sheet must say so."""
		sheet = self.book["Role Profiles"]
		headers = [cell.value for cell in next(sheet.iter_rows(max_row=1))]
		name_col = headers.index("role_profile")
		dupe_col = headers.index("duplicate_of")

		found = {}
		for row in sheet.iter_rows(min_row=2, values_only=True):
			if row[name_col] and row[name_col].startswith("Upande Team"):
				found[row[name_col]] = row[dupe_col] or ""

		self.assertTrue(found, "expected the Upande Team profiles to be present")
		for profile, duplicates in found.items():
			self.assertTrue(duplicates, f"{profile} should list its duplicates")
