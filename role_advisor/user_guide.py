# Copyright (c) 2026, ghost-mann and contributors
# For license information, please see license.txt

"""Build the Role Advisor user guide as a .docx.

Written for the person administering access, not for a developer: tasks and
concepts, no Frappe internals unless a concept is meaningless without one.
"""

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor, Inches

NAVY = RGBColor(0x1F, 0x33, 0x46)
GREY = RGBColor(0x55, 0x5F, 0x66)
RED = RGBColor(0xA3, 0x1D, 0x1D)

doc = Document()

# Base style: a readable serif at a comfortable size beats Calibri 11 for a
# document people will actually sit and read.
normal = doc.styles["Normal"]
normal.font.name = "Georgia"
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(8)
normal.paragraph_format.line_spacing = 1.25

for name, size, colour, before in (
    ("Heading 1", 20, NAVY, 22),
    ("Heading 2", 15, NAVY, 18),
    ("Heading 3", 12, NAVY, 14),
):
    style = doc.styles[name]
    style.font.name = "Georgia"
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = colour
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.keep_with_next = True


def para(text="", style=None, italic=False, colour=None, size=None):
    p = doc.add_paragraph(style=style)
    run = p.add_run(text)
    run.italic = italic
    if colour:
        run.font.color.rgb = colour
    if size:
        run.font.size = Pt(size)
    return p


def rich(parts):
    """A paragraph mixing bold, code and plain runs."""
    p = doc.add_paragraph()
    for text, kind in parts:
        run = p.add_run(text)
        if kind == "b":
            run.bold = True
        elif kind == "c":
            run.font.name = "Consolas"
            run.font.size = Pt(10)
        elif kind == "i":
            run.italic = True
    return p


def bullets(items, style="List Bullet"):
    for item in items:
        if isinstance(item, tuple):
            p = doc.add_paragraph(style=style)
            p.add_run(item[0]).bold = True
            p.add_run(" — " + item[1])
        else:
            doc.add_paragraph(item, style=style)


def _no_split(row):
    """Stop a table row breaking across a page - it reads as truncated."""
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    trPr = row._tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:cantSplit"))


def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(h)
        run.bold = True
        run.font.size = Pt(10)
    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(str(value))
            run.font.size = Pt(10)
    if widths:
        for i, w in enumerate(widths):
            for row in t.rows:
                row.cells[i].width = Inches(w)
    for row in t.rows:
        _no_split(row)
    # Repeat the header if a long table does spill over.
    t.rows[0]._tr.get_or_add_trPr().append(__import__("docx").oxml.OxmlElement("w:tblHeader"))
    doc.add_paragraph()
    return t


def callout(title, body, colour=NAVY):
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    cell = t.rows[0].cells[0]
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(title + "  ")
    run.bold = True
    run.font.color.rgb = colour
    run.font.size = Pt(10)
    run2 = p.add_run(body)
    run2.font.size = Pt(10)
    doc.add_paragraph()


# ───────────────────────────── title ─────────────────────────────
title = para("Role Advisor", "Heading 1")
title.paragraph_format.space_before = Pt(0)
para("A guide for the people who manage who can do what", italic=True, colour=GREY, size=12)
para("Kaitet Group · August 2026 · github.com/ghost-mann/role-advisor", colour=GREY, size=9)

para("What this is for", "Heading 2")
para(
    "Right now, if someone at Karen Roses needs access changed, the request has to go to "
    "Upande. There is no way to let a Karen Roses person handle it, because the only role "
    "that can change access is System Manager — and System Manager can change everything, "
    "for everyone, in every company."
)
para(
    "Role Advisor creates a middle option. A named person can be given the ability to manage "
    "access for their own company only, and to hand out only a specific, pre-approved list of "
    "access levels. Nothing more."
)
rich([
    ("Everything it does is either a report you read, or an assignment you make after seeing "
     "exactly what it will change. ", ""),
    ("It never changes anything silently.", "b"),
])

# ───────────────────────────── concepts ─────────────────────────────
para("The four words you need", "Heading 2")
para(
    "These come from the ERP itself, not from this app, but nothing else makes sense until "
    "they are clear."
)

table(
    ["Word", "What it means", "Example"],
    [
        ["Role", "One permission bundle. The smallest unit. Rarely handed out on its own.",
         "Gate Guard"],
        ["Role Profile", "A named group of roles. This is what a person is actually given.",
         "Agriculture Supervisor"],
        ["Module Profile", "Which parts of the menu a person sees. Tidies navigation only.",
         "Roses Production Supervisor"],
        ["Designation", "The person's job title, from their employee record.",
         "Spray Supervisor"],
    ],
    widths=[1.1, 3.5, 1.6],
)

callout(
    "Important",
    "A Module Profile only hides menu items. It does not stop anyone reaching those screens by "
    "other means. Real access is controlled by the Role Profile. If you want someone genuinely "
    "unable to touch invoices, that is a Role Profile decision, not a Module Profile one.",
    RED,
)

# ───────────────────────────── the console ─────────────────────────────
para("Assigning access", "Heading 2")
para("Open User Access from the sidebar, then User Access Console.")

para("What you see", "Heading 3")
bullets([
    ("The user list", "only the people in your company. Others do not appear, and cannot be "
     "reached even by typing their name."),
    ("Your grantable profiles", "the access levels you personally are allowed to hand out, "
     "listed smallest first so the least-access option is the easiest to pick."),
])

para("How to assign", "Heading 3")
for n, text in enumerate([
    "Search for the person and select them. You will see what they hold today.",
    "Click Preview next to the profile you want to give them.",
    "Read the two lists carefully: what they will gain, and what they will lose.",
    "Confirm, or cancel and pick something else.",
], start=1):
    p = doc.add_paragraph(style="List Number")
    p.add_run(text)

callout(
    "Read the “loses” list",
    "Giving someone a new profile replaces their old one — it does not add to it. Anything the "
    "old profile allowed and the new one does not is removed immediately. The preview shows you "
    "this before you commit, and it is the most common surprise.",
)

para(
    "Every assignment is recorded in Access Assignment Log: who did it, to whom, what changed, "
    "and when. Including assignments that changed nothing. The log cannot be edited or deleted "
    "by anyone."
)

# ───────────────────────────── limits ─────────────────────────────
para("What you cannot do (by design)", "Heading 2")
bullets([
    "Reach anyone outside your company.",
    "Hand out an access level that is not on your list.",
    "Hand out anything that would let the recipient change permissions themselves — refused "
    "even if it were on your list by mistake.",
    "Turn someone into a different kind of user, or touch their password or API keys.",
    "Reach anyone with no employee record. Service accounts and contractors are invisible to "
    "you, deliberately.",
    "Change your own company scope, or widen your own list.",
])
para(
    "Your list of grantable profiles is set by a System Manager on your Delegated User Admin "
    "record. You can see your own record but not change it, and not see anyone else's.",
    italic=True,
)

# ───────────────────────────── map ─────────────────────────────
para("The Designation Access Map", "Heading 2")
para(
    "Rather than deciding access one person at a time forever, the map records the decision "
    "once per job title: a Spray Supervisor at Karen Roses gets Agriculture Supervisor. New "
    "starters can then be assigned in bulk."
)
para(
    "The map has been pre-filled by looking at what people with each job title already hold. "
    "That is evidence, not a recommendation — the existing assignments are known to be "
    "inconsistent, so each row carries a confidence rating and the reasoning behind it."
)

table(
    ["Confidence", "Meaning", "Rows"],
    [
        ["High", "Colleagues in the same job and company overwhelmingly agree.", "4"],
        ["Medium", "Some agreement, but a real split. Needs a decision.", "29"],
        ["Low", "Based on a single colleague. Weak evidence.", "45"],
    ],
    widths=[1.2, 3.8, 0.8],
)

rich([
    ("Only High-confidence rows are switched on. ", "b"),
    ("The other 74 are switched off and will not be used until a human reviews them. Open ", ""),
    ("Designation Access Map", "c"),
    (", filter Is Active to No, and read the Evidence column on each row.", ""),
])

para("Rows worth looking at first", "Heading 3")
bullets([
    ("Irrigator", "currently points at a profile with approval rights. An irrigator should "
     "almost certainly not approve anything."),
    ("Quality Controller", "most colleagues hold a broader profile than the job needs."),
    ("Task Worker", "points at a scouting profile, which looks like a mistake."),
    ("Eight job titles have no precedent at all", "including Executive Director and Stores "
     "Accountant. These need a decision from scratch."),
])

# ───────────────────────────── reports ─────────────────────────────
para("The five reports", "Heading 2")
para("All read-only. None of them changes anything.")

table(
    ["Report", "The question it answers", "What it found"],
    [
        ["Designation Gap", "Who has no access level set at all?",
         "41 people, 17 of whom have no employee record"],
        ["Role Profile Overgrant", "What can each access level actually do?",
         "Upande Team allows almost everything"],
        ["Module Exposure", "Whose menu is untidy?",
         "240 people see every part of the system"],
        ["System Manager Audit", "Who has unrestricted access, and should they?",
         "23 people, plus a service account and two test logins"],
        ["Role Drift", "Is anyone about to lose access unexpectedly?",
         "Nothing at risk"],
    ],
    widths=[1.5, 2.4, 2.3],
)

callout(
    "The one to read first",
    "Role Profile Overgrant. It is the only report that describes what people can actually do, "
    "rather than what they can see. Everything else is context.",
)

# ───────────────────────────── findings ─────────────────────────────
para("What the review has already turned up", "Heading 2")
para("Findings from the existing setup, in rough order of how much they matter.")

bullets([
    ("Almost everyone can reach almost everything", "316 of roughly 333 people with an access "
     "level can reach Accounts; 309 can reach HR. A handful of modules are properly scoped; the "
     "core business ones are wide open."),
    ("Security Guard grants nothing", "16 people hold an access level that permits no action at "
     "all. Either it is incomplete, or those people need a different one."),
    ("Three access levels are identical", "Upande Team, Upande Team - COO and Upande Team HR "
     "Users allow exactly the same things. The different names promise a difference that does "
     "not exist."),
    ("Temporary staff have more access than permanent", "Temporary -Mechanic allows considerably "
     "more than Mechanic."),
    ("Two test logins have unrestricted access to live", "eric@test.com and test@upande.com both "
     "hold System Manager on the production system, alongside a service account."),
    ("66 disabled accounts still hold access levels", "They cannot log in today, but every "
     "permission is intact. Re-enable one and it silently gets everything back."),
    ("Naming is inconsistent", "40 of 87 access levels have a naming problem — mixed "
     "capitalisation, inconsistent punctuation, and place names mixed with authority levels in "
     "the same position."),
])

# ───────────────────────────── glossary of screens ─────────────────────────────
para("Where everything lives", "Heading 2")
table(
    ["Screen", "Use it to"],
    [
        ["User Access Console", "Assign access to one person, with a preview first"],
        ["Designation Access Map", "Record and review the access level for each job title"],
        ["Delegated User Admin", "See your own scope and grantable list (read-only for you)"],
        ["Access Assignment Log", "See every change ever made, by whom"],
        ["User Access Settings", "System Manager only. Thresholds and policy."],
        ["The five reports", "Understand the current state before changing it"],
    ],
    widths=[2.2, 4.0],
)

para("If something is refused", "Heading 2")
para(
    "The app explains why rather than failing silently. The common ones:"
)
bullets([
    ("“You are not configured as a delegated user administrator”", "you have no Delegated User "
     "Admin record, or it is switched off. A System Manager sets this up."),
    ("“… is outside the users you administer”", "that person belongs to another company, or has "
     "no employee record."),
    ("“… is not one of the role profiles you may grant”", "ask for it to be added to your list, "
     "if it is appropriate."),
    ("“… cannot be delegated”", "that access level would let the recipient change permissions. "
     "It can only be given by a System Manager."),
])

para("A note on the current stage", "Heading 2")
para(
    "The app is installed and working, but no restrictions apply to anyone until a Delegated "
    "User Admin record is created for them. Until then, everything behaves exactly as it did "
    "before. The first two people intended for this are the Karen Roses IT staff, who already "
    "hold the User Manager role and currently cannot complete the task it implies.",
)

def build(path: str | None = None) -> str:
    """Regenerate the guide. Run after any change to the workflow it describes."""
    import os

    if not path:
        app_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
        path = os.path.join(app_dir, "docs", "Role Advisor - User Guide.docx")

    doc.save(path)
    print(path)

    return path


if __name__ == "__main__":
    build()
