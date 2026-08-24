# Copyright (c) 2026, ghost-mann and contributors
# For license information, please see license.txt

import os

import frappe
import openpyxl
from frappe.tests import IntegrationTestCase

from upande_core import capability, module_workbook, workbook

FRONT_TABS = [
	"Overview",
	"Users",
	"Users by Module",
	"Role Profiles",
	"Roles",
	"Profile Index",
	"Permissions",
]


class TestCombinedWorkbook(IntegrationTestCase):
	@classmethod
	def setUpClass(cls):
		super().setUpClass()
		capability.clear_capability_index()
		cls.path = workbook.build_combined(
			path=os.path.join(
				frappe.utils.get_site_path("private", "files"), "_ra_test_combined.xlsx"
			)
		)
		cls.book = openpyxl.load_workbook(cls.path, read_only=True)

	@classmethod
	def tearDownClass(cls):
		cls.book.close()
		if os.path.exists(cls.path):
			os.remove(cls.path)
		super().tearDownClass()

	def test_every_sheet_from_both_workbooks_is_present(self):
		modules = len(module_workbook.gather()["by_module"])

		self.assertEqual(self.book.sheetnames[: len(FRONT_TABS)], FRONT_TABS)
		self.assertEqual(len(self.book.sheetnames), len(FRONT_TABS) + modules)

	def test_tab_names_stay_unique_against_the_front_tabs(self):
		"""A module called 'Users' must not collide with the Users sheet."""
		self.assertEqual(len(self.book.sheetnames), len(set(self.book.sheetnames)))
		for name in self.book.sheetnames:
			self.assertLessEqual(len(name), 31, name)

	def test_users_sheet_holds_every_user(self):
		self.assertEqual(self.book["Users"].max_row - 1, frappe.db.count("User"))

	def test_disabled_users_reach_the_module_sheets(self):
		"""They hold profiles, so they appear against the modules those reach."""
		sheet = self.book["Users by Module"]
		headers = [cell.value for cell in next(sheet.iter_rows(max_row=1))]
		column = headers.index("enabled")

		values = {row[column] for row in sheet.iter_rows(min_row=2, values_only=True)}

		self.assertIn(0, values)

	def test_overview_separates_enabled_from_total(self):
		sheet = self.book["Overview"]
		headers = [cell.value for cell in next(sheet.iter_rows(max_row=1))]

		# Both counts are needed: a module reachable by 383 accounts of which
		# only 317 can log in is a different picture from 383 live users.
		self.assertIn("users_with_access", headers)
		self.assertIn("enabled_with_access", headers)

		total = headers.index("users_with_access")
		enabled = headers.index("enabled_with_access")
		for row in sheet.iter_rows(min_row=2, values_only=True):
			self.assertGreaterEqual(row[total], row[enabled])

	def test_the_export_writes_nothing_to_the_site(self):
		before = (frappe.db.count("User"), frappe.db.count("Role Profile"))

		second = os.path.join(
			frappe.utils.get_site_path("private", "files"), "_ra_test_combined2.xlsx"
		)
		workbook.build_combined(path=second)
		os.remove(second)

		self.assertEqual(
			before, (frappe.db.count("User"), frappe.db.count("Role Profile"))
		)
